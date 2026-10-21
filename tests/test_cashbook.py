from contextlib import closing
from pathlib import Path
import sqlite3
import tempfile
import unittest

from finance_ledger.categories import CATEGORY_REPORTS, TRANSACTION_DETAILS
from finance_ledger.ledger import Ledger, TYPES
from finance_ledger.reports import CASHBOOK_REPORT, build_report, cashbook_balances, export_report
from finance_ledger.schema import APPLICATION_ID, MIGRATIONS
from finance_ledger.storage import Database


class CashbookCase(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.db=Database(self.temp.name,daily=False)
        self.ledger=Ledger(self.db)
        self.cash=self.ledger.add_cash('Cash','cash',10_000_000,'2026-01-01')
        self.request=dict(type_key='loan_disbursed',amount_paise=1_000_000,tx_date='2026-01-02',cash_id=self.cash,idempotency_key='intake-1',new_party={'name':'New borrower','phone':'1234567890'},new_finance={'reference':'STL-001','annual_rate':'12'},particulars='Business loan')

    def tearDown(self):
        self.db.close(); self.temp.cleanup()

    def test_inline_party_account_and_posting_are_atomic(self):
        group=self.ledger.post_cashbook(**self.request)
        entry=self.ledger.entries()[0]
        self.assertEqual(entry['group_id'],group)
        self.assertEqual(entry['particulars'],'Business loan')
        self.assertEqual(entry['party_name'],'New borrower')
        self.assertEqual(entry['finance_reference'],'STL-001')
        self.assertEqual(self.ledger.principal(entry['finance_id'],'2026-01-31'),1_000_000)
        self.assertEqual(self.ledger.cash_balance(self.cash,'2026-01-31'),9_000_000)
        self.assertEqual(self.db.one('SELECT COUNT(*) n FROM audit_log')['n'],4)

    def test_failed_posting_rolls_back_party_account_and_audits(self):
        request=self.request | {'type_key':'loan_repaid'}
        with self.assertRaisesRegex(ValueError,'negative'):
            self.ledger.post_cashbook(**request)
        for table in ('parties','finance_accounts','posting_groups','entries'):
            self.assertEqual(self.db.one(f'SELECT COUNT(*) n FROM {table}')['n'],0)
        self.assertEqual(self.db.one('SELECT COUNT(*) n FROM audit_log')['n'],1)
        # The same form can be corrected and submitted using its original key.
        self.ledger.post_cashbook(**self.request)
        self.assertEqual(len(self.ledger.entries()),1)

    def test_retry_does_not_create_duplicate_names_or_accounts(self):
        group=self.ledger.post_cashbook(**self.request)
        self.assertEqual(self.ledger.post_cashbook(**self.request),group)
        self.assertEqual(len(self.ledger.parties()),1)
        self.assertEqual(len(self.ledger.finance_accounts()),1)
        self.assertEqual(len(self.ledger.entries()),1)
        with self.assertRaisesRegex(ValueError,'different values'):
            self.ledger.post_cashbook(**(self.request | {'particulars':'Altered'}))
        self.assertEqual(len(self.ledger.parties()),1)

    def test_existing_party_new_account_and_split_receipt(self):
        party=self.ledger.save_party({'name':'Existing borrower'})
        request=self.request | {'type_key':'loan_repaid','new_party':None,'party_id':party,'new_finance':{'reference':'OLD-STL','opening_paise':1_000_000},'amount_paise':200_000,'interest_paise':2000}
        group=self.ledger.post_cashbook(**request)
        entries=self.ledger.entries()
        self.assertEqual(len(entries),2)
        self.assertEqual({r['group_id'] for r in entries},{group})
        self.assertEqual(self.ledger.principal(entries[0]['finance_id'],'2026-01-31'),800_000)
        self.assertEqual(self.ledger.dashboard('2026-01-01','2026-01-31')['interest_received'],2000)
        self.ledger.reverse(group,'2026-01-03','Wrong receipt','reversal')
        self.assertTrue(all(r['particulars']=='Business loan' for r in self.ledger.entries()))
        self.assertEqual(self.ledger.principal(entries[0]['finance_id'],'2026-01-31'),1_000_000)

    def test_office_particulars_are_not_fictitious_parties(self):
        self.ledger.post_cashbook(type_key='general_expense',amount_paise=20_000,tx_date='2026-01-02',cash_id=self.cash,idempotency_key='expense',particulars='Water Can')
        self.assertEqual(self.ledger.parties(),[])
        row=self.ledger.entries()[0]
        self.assertEqual(row['particulars'],'Water Can')
        self.assertIsNone(row['party_id'])
        self.assertEqual(row['direction'],-1)
        report=build_report(self.ledger,CATEGORY_REPORTS['general_expense'],'2026-01-01','2026-01-31')
        self.assertEqual(report.rows[1][3:5],['','Water Can'])
        self.assertEqual(report.rows[-1][6],20_000)
        for suffix in ('.csv','.xlsx','.pdf'):
            export_report(report,Path(self.temp.name)/('office'+suffix))

    def test_reject_ambiguous_new_and_existing_selections(self):
        self.ledger.post_cashbook(**self.request)
        entry=self.ledger.entries()[0]
        variants=[{'party_id':entry['party_id']},{'finance_id':entry['finance_id']},{'new_finance':{'module':'fd'}},{'type_key':'general_income'},{'interest_paise':100}]
        for changes in variants:
            with self.subTest(changes=changes),self.assertRaises(ValueError):
                self.ledger.post_cashbook(**(self.request | {'idempotency_key':'invalid'} | changes))
        self.assertEqual(len(self.ledger.entries()),1)

    def test_duplicate_party_name_requires_existing_selection(self):
        self.ledger.save_party({'name':'New borrower'})
        with self.assertRaisesRegex(ValueError,'already exists'):
            self.ledger.post_cashbook(**self.request)
        self.assertEqual(len(self.ledger.parties()),1)
        self.assertEqual(self.ledger.finance_accounts(),[])

    def test_all_finance_families_use_correct_accounting(self):
        types=['fd_received','loan_disbursed','other_received','hand_given_disbursed','hand_taken_received']
        for index,key in enumerate(types):
            with self.subTest(type=key):
                self.ledger.post_cashbook(**(self.request | {'type_key':key,'idempotency_key':key,'new_party':{'name':key},'new_finance':{}}))
                row=self.ledger.entries()[-1]
                self.assertEqual(row['direction'],TYPES[key].direction)
                self.assertEqual(row['module'],TYPES[key].module)
                self.assertTrue(row['finance_reference'])
        self.assertEqual(len(self.ledger.finance_accounts()),5)

    def test_combined_reports_include_debit_credit_interest_without_double_counting(self):
        self.ledger.post_cashbook(**self.request)
        account=self.ledger.entries()[0]['finance_id']
        self.ledger.post_cashbook(type_key='loan_repaid',amount_paise=100_000,interest_paise=1000,tx_date='2026-01-03',cash_id=self.cash,idempotency_key='receipt',finance_id=account,particulars='Receipt')
        self.ledger.post_cashbook(type_key='general_expense',amount_paise=500,tx_date='2026-01-03',cash_id=self.cash,idempotency_key='office',particulars='Water')
        report=build_report(self.ledger,CATEGORY_REPORTS['stl'],'2026-01-01','2026-01-31',type='loan_repaid',module='fd')
        self.assertEqual(len(report.rows),5)
        self.assertEqual(report.rows[-1][5:9],[101_000,1_000_000,-899_000,1000])
        self.assertEqual(report.rows[2][8],0)  # principal repayment
        self.assertEqual(report.rows[3][8],1000)  # interest, already in cash in
        feb=build_report(self.ledger,CATEGORY_REPORTS['stl'],'2026-02-01','2026-02-28')
        self.assertEqual(feb.rows[0][7],-899_000)
        self.assertEqual(feb.rows[-1][7],-899_000)
        for key,name in CATEGORY_REPORTS.items():
            result=build_report(self.ledger,name,'2026-01-01','2026-01-31')
            self.assertTrue(all(len(r)==len(result.columns) for r in result.rows))
        self.assertEqual(len(CATEGORY_REPORTS),8)
        self.assertEqual(set(TRANSACTION_DETAILS),set(TYPES)-{'transfer_in'})

    def test_excel_cashbook_has_typed_dates_money_and_reference_layout(self):
        from datetime import datetime
        from openpyxl import load_workbook
        self.ledger.post_cashbook(**self.request)
        report=build_report(self.ledger,CATEGORY_REPORTS['stl'],'2026-01-01','2026-01-31')
        path=Path(self.temp.name)/'combined.xlsx'
        export_report(report,path)
        workbook=load_workbook(path); sheet=workbook.active
        self.assertEqual([sheet.cell(4,i).value for i in range(1,12)],report.columns)
        self.assertEqual(sheet['G6'].value,10000)
        self.assertEqual(sheet['H7'].value,-10000)
        self.assertIsInstance(sheet['C6'].value,datetime)
        self.assertEqual(sheet['C6'].number_format,'dd.mm.yyyy')
        self.assertEqual(sheet.freeze_panes,'F6')
        self.assertEqual(sheet['F6'].border.bottom.style,'thin')
        workbook.close()

    def test_each_finance_report_keeps_interest_direction_and_reversals(self):
        families=[('stl','loan_disbursed','loan_repaid','loan_interest',1),
                  ('fd','fd_received','fd_repaid','fd_interest',-1),
                  ('external','other_received','other_repaid','other_interest',-1),
                  ('hand_given','hand_given_disbursed','hand_given_repaid','hand_given_interest',1),
                  ('hand_taken','hand_taken_received','hand_taken_repaid','hand_taken_interest',-1)]
        for key,initial,repayment,interest,direction in families:
            with self.subTest(family=key):
                self.ledger.post_cashbook(**(self.request | {'type_key':initial,'idempotency_key':key,'new_party':{'name':key},'new_finance':{'reference':key}}))
                account=next(a['id'] for a in self.ledger.finance_accounts() if a['reference']==key)
                group=self.ledger.post_cashbook(type_key=repayment,amount_paise=100_000,tx_date='2026-01-03',cash_id=self.cash,idempotency_key=key+'-repay',finance_id=account)
                interest_group=self.ledger.post_cashbook(type_key=interest,amount_paise=1000,tx_date='2026-01-03',cash_id=self.cash,idempotency_key=key+'-interest',finance_id=account)
                self.ledger.reverse(group,'2026-02-01','Correction',key+'-reverse')
                self.ledger.reverse(interest_group,'2026-02-01','Correction',key+'-reverse-interest')
                jan=build_report(self.ledger,CATEGORY_REPORTS[key],'2026-01-01','2026-01-31')
                feb=build_report(self.ledger,CATEGORY_REPORTS[key],'2026-02-01','2026-02-28')
                self.assertEqual(len(jan.rows),5)
                self.assertEqual({r[1] for r in jan.rows[1:-1]}, {TRANSACTION_DETAILS[t] for t in (initial,repayment,interest)})
                self.assertEqual(jan.rows[-1][8],1000*direction)
                self.assertEqual(feb.rows[-1][8],-1000*direction)
                self.assertEqual(feb.rows[0][7],jan.rows[-1][7])
                self.assertEqual(feb.rows[-1][7],-1_000_000*direction)
                scoped=build_report(self.ledger,CATEGORY_REPORTS[key],'2026-01-01','2026-01-31',finance_id=account)
                self.assertEqual(scoped.rows,jan.rows)

    def test_running_cash_stays_chronological_and_includes_opening(self):
        self.ledger.post_cashbook(**self.request)
        self.ledger.post_cashbook(type_key='general_expense',amount_paise=1000,tx_date='2026-01-03',cash_id=self.cash,idempotency_key='rent',particulars='Rent')
        report=build_report(self.ledger,CASHBOOK_REPORT,'2026-01-01','2026-01-31')
        self.assertEqual(report.rows[-1][7],8_999_000)
        self.assertEqual(report.rows[-1][6],1_001_000)
        opening,closing,balances=cashbook_balances(self.ledger,'2026-01-03','2026-01-31')
        self.assertEqual((opening,closing),(9_000_000,8_999_000))
        self.assertEqual(balances[self.ledger.entries()[-1]['id']],8_999_000)
        historical=build_report(self.ledger,CASHBOOK_REPORT,'2026-01-01','2026-01-02')
        self.assertEqual(historical.rows[-1][7],9_000_000)

    def test_category_report_reversals_do_not_rewrite_previous_period(self):
        group=self.ledger.post_cashbook(**self.request)
        self.ledger.reverse(group,'2026-02-01','Correction','reverse')
        january=build_report(self.ledger,CATEGORY_REPORTS['stl'],'2026-01-01','2026-01-31')
        february=build_report(self.ledger,CATEGORY_REPORTS['stl'],'2026-02-01','2026-02-28')
        self.assertEqual(january.rows[-1][6],1_000_000)
        self.assertEqual(february.rows[-1][6],-1_000_000)
        self.assertEqual(february.rows[1][4],'Business loan')


class VersionOneMigrationCase(unittest.TestCase):
    def test_old_database_and_backup_upgrade_preserve_every_posting(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'ledger.sqlite3'
            with closing(sqlite3.connect(path)) as conn:
                conn.executescript(MIGRATIONS[0][1])
                conn.execute(f'PRAGMA application_id={APPLICATION_ID}')
                conn.execute('PRAGMA user_version=1')
                conn.execute("INSERT INTO cash_accounts(name,kind,opening_paise,opening_date) VALUES('Cash','cash',100000,'2026-01-01')")
                conn.execute("INSERT INTO parties(name,created_at,updated_at) VALUES('Old party','2026-01-01T00:00:00+00:00','2026-01-01T00:00:00+00:00')")
                conn.execute("INSERT INTO finance_accounts(reference,party_id,module,start_date,opening_paise,created_at) VALUES('OLD-FD',1,'fd','2026-01-01',0,'2026-01-01T00:00:00+00:00')")
                conn.execute("INSERT INTO posting_groups(id,reference,idempotency_key,fingerprint,tx_date,method,created_by,created_at) VALUES('old','TX-OLD','old-key','old-fingerprint','2026-01-02','cash','Old user','2026-01-02T00:00:00+00:00')")
                conn.execute("INSERT INTO entries(group_id,cash_id,party_id,finance_id,type,amount_paise,direction,principal_delta,interest,factor) VALUES('old',1,1,1,'fd_received',50000,1,1,0,1)")
                conn.commit()
            db=Database(folder,daily=False)
            try:
                ledger=Ledger(db)
                self.assertEqual(db.conn.execute('PRAGMA user_version').fetchone()[0],2)
                self.assertEqual(ledger.cash_balance(1,'2026-01-31'),150000)
                self.assertEqual(ledger.principal(1,'2026-01-31'),50000)
                self.assertEqual(ledger.entries()[0]['particulars'],'')
                backups=list((Path(folder)/'backups').glob('pre-migration-*.sqlite3'))
                self.assertEqual(len(backups),1)
                db.validate_backup(backups[0])
                ledger.post_cashbook(type_key='fd_interest',amount_paise=100,tx_date='2026-01-03',cash_id=1,finance_id=1,idempotency_key='new',particulars='Interest')
                self.assertEqual(len(ledger.entries()),2)
                db.restore(backups[0])
                self.assertEqual(len(ledger.entries()),1)
                self.assertEqual(ledger.principal(1,'2026-01-31'),50000)
                self.assertEqual(db.conn.execute('PRAGMA user_version').fetchone()[0],2)
            finally:
                db.close()
