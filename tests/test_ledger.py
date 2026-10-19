from contextlib import closing
import csv
import json
import sqlite3
import subprocess
import sys
import tempfile
import unittest
import uuid
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

from finance_ledger.demo import seed_demo
from finance_ledger.ledger import Ledger, TYPES
from finance_ledger.money import interest_paise, money, paise, today
from finance_ledger.reports import REPORT_NAMES, Report, build_report, export_report
from finance_ledger.schema import APPLICATION_ID, MIGRATIONS
from finance_ledger.storage import AlreadyRunning, Database


class LedgerCase(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.db=Database(self.temp.name,daily=False)
        self.l=Ledger(self.db)
        self.cash=self.l.add_cash('Cash','cash',10_000_000,'2026-01-01')
        self.bank=self.l.add_cash('Bank','bank',0,'2026-01-01')
        self.party=self.l.save_party({'name':'Test Party','roles':'FD depositor, Loan borrower'})
        self.accounts={m:self.l.add_finance(m.upper(),self.party,m,'2026-01-01',annual_rate='10') for m in ('fd','loan','other','hand_given','hand_taken')}

    def tearDown(self):
        self.db.close(); self.temp.cleanup()

    def post(self,type_key,amount,day='2026-01-02',**kwargs):
        return self.l.post(tx_date=day,components=[(type_key,amount)],cash_id=self.cash,idempotency_key=kwargs.pop('idempotency_key',str(uuid.uuid4())),finance_id=kwargs.pop('finance_id',self.accounts.get(TYPES[type_key].module)),**kwargs)

    def test_required_example(self):
        self.post('fd_received',5_000_000)
        self.post('loan_disbursed',2_000_000)
        self.l.post(tx_date='2026-01-03',components=[('loan_repaid',500_000),('loan_interest',50_000)],cash_id=self.cash,finance_id=self.accounts['loan'],idempotency_key='split')
        self.post('fd_interest',30_000,'2026-01-04')
        result=self.l.dashboard('2026-01-01','2026-01-31')
        self.assertEqual(result['cash'],13_520_000)
        self.assertEqual(result['principal']['fd'],5_000_000)
        self.assertEqual(result['principal']['loan'],1_500_000)
        self.assertEqual(result['interest_received'],50_000)
        self.assertEqual(result['interest_paid'],30_000)
        self.assertEqual(result['net_interest'],20_000)
        self.assertEqual(result['amount_in'],5_550_000)
        self.assertEqual(result['amount_out'],2_030_000)

    def test_every_type_derives_business_direction(self):
        expected={'fd_received':1,'fd_repaid':-1,'fd_interest':-1,'loan_disbursed':-1,'loan_repaid':1,'loan_interest':1,'other_received':1,'other_repaid':-1,'other_interest':-1,'hand_given_disbursed':-1,'hand_given_repaid':1,'hand_given_interest':1,'hand_taken_received':1,'hand_taken_repaid':-1,'hand_taken_interest':-1,'general_income':1,'general_expense':-1,'transfer_in':1,'transfer_out':-1}
        self.assertEqual({k:v.direction for k,v in TYPES.items()},expected)
        # Post each finance type, including all liability and receivable families.
        for key,spec in TYPES.items():
            if not spec.module: continue
            self.post(key,1000)
            row=self.l.entries()[-1]
            self.assertEqual(row['direction'],expected[key])
            self.assertEqual(row['principal_delta'],spec.principal)

    def test_partial_repayments_and_as_of(self):
        self.post('loan_disbursed',100_000,'2026-01-02')
        self.post('loan_repaid',25_000,'2026-01-10')
        self.post('loan_repaid',10_000,'2026-01-05')
        self.assertEqual(self.l.principal(self.accounts['loan'],'2026-01-04'),100_000)
        self.assertEqual(self.l.principal(self.accounts['loan'],'2026-01-06'),90_000)
        self.assertEqual(self.l.principal(self.accounts['loan'],'2026-01-31'),65_000)
        report=build_report(self.l,'Principal outstanding as of date','2026-01-01','2026-01-06')
        row=next(r for r in report.rows if r[0]=='LOAN')
        self.assertEqual(row[9],90_000)

    def test_backdated_posting_must_not_break_future_balance(self):
        self.post('loan_disbursed',100_000,'2026-01-02')
        self.post('loan_repaid',100_000,'2026-01-20')
        count=len(self.l.entries())
        with self.assertRaisesRegex(ValueError,'negative on 2026-01-20'):
            self.post('loan_repaid',1,'2026-01-10')
        self.assertEqual(len(self.l.entries()),count)

    def test_split_is_atomic(self):
        self.post('loan_disbursed',100_000)
        with self.assertRaises(ValueError):
            self.l.post(tx_date='2026-01-03',components=[('loan_interest',1000),('loan_repaid',100_001)],cash_id=self.cash,finance_id=self.accounts['loan'],idempotency_key='bad')
        self.assertEqual(len(self.l.entries()),1)
        self.assertIsNone(self.db.one('SELECT id FROM posting_groups WHERE idempotency_key=?',('bad',)))

    def test_split_reversal_is_atomic_and_historical(self):
        self.post('loan_disbursed',100_000)
        group=self.l.post(tx_date='2026-01-03',components=[('loan_repaid',25_000),('loan_interest',2500)],cash_id=self.cash,finance_id=self.accounts['loan'],idempotency_key='split')
        reversal=self.l.reverse(group,'2026-02-01','Wrong receipt','reverse')
        self.assertEqual(self.l.principal(self.accounts['loan'],'2026-01-31'),75_000)
        self.assertEqual(self.l.principal(self.accounts['loan'],'2026-02-01'),100_000)
        self.assertEqual(len([r for r in self.l.entries() if r['group_id']==reversal]),2)
        self.assertEqual(self.l.dashboard('2026-01-01','2026-02-01')['interest_received'],0)
        self.assertEqual(self.l.dashboard('2026-02-01','2026-02-01')['amount_in'],-27_500)
        self.assertEqual(self.l.reverse(group,'2026-02-01','Wrong receipt','reverse'),reversal)
        with self.assertRaises(ValueError): self.l.reverse(group,'2026-02-02','Again','different')
        with self.assertRaises(ValueError): self.l.reverse(reversal,'2026-02-02','Again','third')

    def test_cannot_reverse_disbursement_under_repayments(self):
        group=self.post('loan_disbursed',100_000)
        self.post('loan_repaid',25_000,'2026-01-03')
        with self.assertRaises(ValueError): self.l.reverse(group,'2026-01-04','Wrong','reverse')
        self.assertEqual(len(self.l.entries()),2)

    def test_reversal_date_and_reason(self):
        group=self.post('fd_received',1000)
        with self.assertRaises(ValueError): self.l.reverse(group,'2026-01-01','Wrong','reverse')
        with self.assertRaises(ValueError): self.l.reverse(group,'2026-01-03','','reverse')

    def test_duplicate_submissions(self):
        group=self.post('fd_received',1000,idempotency_key='unique')
        self.assertEqual(self.post('fd_received',1000,idempotency_key='unique'),group)
        self.assertEqual(len(self.l.entries()),1)
        with self.assertRaises(ValueError): self.post('fd_received',2000,idempotency_key='unique')

    def test_duplicate_key_is_database_constrained(self):
        self.post('fd_received',1000,idempotency_key='unique')
        with self.assertRaises(sqlite3.IntegrityError):
            self.db.conn.execute("INSERT INTO posting_groups SELECT 'other-id','OTHER-REF',idempotency_key,fingerprint,tx_date,method,external_ref,notes,NULL,created_by,created_at FROM posting_groups")

    def test_transfer_atomic_and_excluded_from_overall(self):
        group=self.post('transfer_out',50_000,transfer_to=self.bank)
        self.assertEqual(len(self.l.entries()),2)
        self.assertEqual(self.l.cash_balance(self.cash,'2026-01-31'),9_950_000)
        self.assertEqual(self.l.cash_balance(self.bank,'2026-01-31'),50_000)
        result=self.l.dashboard('2026-01-01','2026-01-31')
        self.assertEqual((result['amount_in'],result['amount_out'],result['cash']),(0,0,10_000_000))
        report=build_report(self.l,'Cash book / Bank ledger','2026-01-01','2026-01-31',cash_id=self.bank)
        self.assertEqual(report.rows[-1][-1],50_000)
        self.l.reverse(group,'2026-01-03','Wrong transfer','reverse')
        self.assertEqual(self.l.cash_balance(self.bank,'2026-01-31'),0)
        self.assertEqual(self.l.cash_balance(self.cash,'2026-01-31'),10_000_000)

    def test_transfer_invalid_destination_rolls_back(self):
        for target in (self.cash,999):
            with self.assertRaises(ValueError): self.post('transfer_out',1000,transfer_to=target)
        self.assertFalse(self.l.entries())

    def test_opening_principal_and_cash_not_income(self):
        account=self.l.add_finance('OLD',self.party,'fd','2026-01-10',50_000)
        self.assertEqual(self.l.principal(account,'2026-01-09'),0)
        self.assertEqual(self.l.principal(account,'2026-01-10'),50_000)
        self.assertEqual(self.l.dashboard('2026-01-01','2026-01-31')['amount_in'],0)
        report=build_report(self.l,'Individual FD statement','2026-01-01','2026-01-31',finance_id=account)
        self.assertEqual(report.rows[-1][-1],50_000)
        report=build_report(self.l,'Cash book / Bank ledger','2026-01-01','2026-01-31',cash_id=self.cash)
        self.assertEqual(report.rows[-1][-1],10_000_000)
        self.assertEqual(report.rows[-1][6],0)

    def test_closed_account_and_explicit_reopen(self):
        account=self.accounts['loan']
        self.l.update_finance(account,'closed','10',None,'Monthly','')
        with self.assertRaises(ValueError): self.post('loan_disbursed',1000)
        self.l.update_finance(account,'active','10',None,'Monthly','Reopened')
        self.post('loan_disbursed',1000)
        with self.assertRaises(ValueError): self.l.update_finance(account,'closed','10',None,'Monthly','')

    def test_types_and_dates_validated(self):
        for value in (0,-1,1.1,True,10**17):
            with self.assertRaises(ValueError): self.post('fd_received',value)
        with self.assertRaises(ValueError): self.post('loan_disbursed',1000,finance_id=self.accounts['fd'])
        with self.assertRaises(ValueError): self.post('loan_disbursed',1000,'2025-12-31')
        with self.assertRaises(ValueError): self.post('loan_disbursed',1000,'2026-02-30')
        with self.assertRaises(ValueError): self.post('loan_disbursed',1000,party_id=999)

    def test_general_income_expense_categories(self):
        income=self.l.add_category('Service fees','income')
        expense=self.l.add_category('Rent','expense')
        self.post('general_income',1000,category_id=income)
        self.post('general_expense',500,category_id=expense)
        with self.assertRaises(ValueError): self.post('general_income',1000,category_id=expense)
        result=self.l.dashboard('2026-01-01','2026-01-31')
        self.assertEqual((result['general_income'],result['general_expense']),(1000,500))
        self.assertEqual(result['net_interest'],0)

    def test_actual_days_estimate_accounts_for_partial_repayment(self):
        self.post('loan_disbursed',100_000,'2026-01-01')
        self.post('loan_repaid',50_000,'2026-01-11')
        result=self.l.estimate(self.accounts['loan'],'2026-01-01','2026-01-21')
        self.assertEqual(result['principal_days'],1_500_000)
        self.assertEqual(result['paise'],411)
        self.assertEqual(len(self.l.entries()),2)
        self.assertEqual(self.l.dashboard('2026-01-01','2026-01-31')['interest_received'],0)

    def test_interest_free_hand_loans(self):
        account=self.accounts['hand_given']
        self.l.update_finance(account,'active','0',None,'Interest-free','')
        self.post('hand_given_disbursed',10_000)
        self.assertEqual(self.l.estimate(account,'2026-01-02','2026-02-02')['paise'],0)

    def test_reports_all_modules_and_period_opening(self):
        self.post('fd_received',100_000)
        self.post('loan_disbursed',100_000)
        self.post('other_received',100_000)
        self.post('hand_given_disbursed',100_000)
        self.post('hand_taken_received',100_000)
        self.post('loan_repaid',20_000,'2026-02-01')
        report=build_report(self.l,'Individual loan statement','2026-02-01','2026-02-28',finance_id=self.accounts['loan'])
        self.assertEqual(report.rows[0][-1],100_000)
        self.assertEqual(report.rows[-1][-1],80_000)
        for name in REPORT_NAMES:
            kwargs={}
            if name=='Cash book / Bank ledger': kwargs['cash_id']=self.cash
            if name=='Party statement': kwargs['party_id']=self.party
            if name=='Individual FD statement': kwargs['finance_id']=self.accounts['fd']
            if name=='Individual loan statement': kwargs['finance_id']=self.accounts['loan']
            if name=='Other finance statement': kwargs['finance_id']=self.accounts['other']
            with self.subTest(name=name):
                report=build_report(self.l,name,'2026-01-01','2026-02-28',**kwargs)
                self.assertTrue(report.columns)
                self.assertTrue(all(len(r)==len(report.columns) for r in report.rows))
                export_report(report,Path(self.temp.name)/'report.csv')

    def test_filters_and_historical_party_totals(self):
        self.post('fd_received',1000)
        self.post('fd_interest',10,'2026-02-01')
        rows=self.l.entries('2026-01-01','2026-01-31',party_id=self.party,module='fd',cash_id=self.cash,type='fd_received')
        self.assertEqual(len(rows),1)
        report=build_report(self.l,'Party statement','2026-01-01','2026-01-31',party_id=self.party)
        self.assertTrue(any('Actual interest In / Out: Rs 0.00 / Rs 0.00' in note for note in report.notes))

    def test_overall_and_monthly_reconcile_to_actual_cash(self):
        self.post('fd_received',5_000_000)
        self.post('loan_disbursed',2_000_000)
        self.post('loan_repaid',500_000,'2026-01-10')
        self.post('loan_interest',50_000,'2026-01-10')
        self.post('fd_interest',30_000,'2026-01-15')
        self.post('transfer_out',100_000,'2026-01-16',transfer_to=self.bank)
        report=build_report(self.l,'Overall Amount In / Amount Out','2026-01-01','2026-01-31')
        self.assertEqual(report.rows[-1][-1],13_520_000)
        self.assertEqual(report.rows[-1][6:8],[5_550_000,2_030_000])
        monthly=build_report(self.l,'Monthly summary','2026-01-01','2026-02-28')
        self.assertEqual(monthly.rows[0][1:7],[0,10_000_000,5_550_000,2_030_000,3_520_000,13_520_000])
        self.assertEqual(monthly.rows[1][1],13_520_000)
        self.assertEqual(monthly.rows[1][6],13_520_000)
        report=build_report(self.l,'Overall Amount In / Amount Out','2026-01-01','2026-01-31',cash_id=self.bank)
        self.assertEqual(report.rows[-1][-1],100_000)
        self.assertEqual(report.rows[-1][6:8],[0,0])

    def test_failed_restore_reinstates_pre_restore_records(self):
        self.post('fd_received',1000)
        path=Path(self.temp.name)/'restore.sqlite3'; self.db.backup(path)
        self.post('fd_received',2000)
        with patch.object(self.db,'_migrate',side_effect=ValueError('Simulated upgrade failure')):
            with self.assertRaisesRegex(ValueError,'Simulated'):
                self.db.restore(path)
        self.assertEqual(self.l.principal(self.accounts['fd'],'2026-01-31'),3000)
        self.assertEqual(self.db.conn.execute('PRAGMA foreign_keys').fetchone()[0],1)

    def test_no_deletion_or_mutation_of_financial_history(self):
        self.post('fd_received',1000)
        for sql in ('DELETE FROM entries','UPDATE entries SET amount_paise=1','DELETE FROM posting_groups','UPDATE posting_groups SET notes=\'changed\'','DELETE FROM audit_log','UPDATE audit_log SET action=\'changed\''):
            with self.subTest(sql=sql), self.assertRaises(sqlite3.IntegrityError): self.db.conn.execute(sql)
        self.assertEqual(self.db.conn.execute('PRAGMA foreign_keys').fetchone()[0],1)

    def test_foreign_keys_and_positive_integer_constraints(self):
        self.post('fd_received',1000)
        with self.assertRaises(sqlite3.IntegrityError): self.db.conn.execute('DELETE FROM parties WHERE id=?',(self.party,))
        for amount in (0,-1,0.5):
            with self.assertRaises(sqlite3.IntegrityError):
                self.db.conn.execute('INSERT INTO entries(group_id,cash_id,type,amount_paise,direction,principal_delta,interest,factor) SELECT group_id,cash_id,type,?,direction,principal_delta,interest,factor FROM entries LIMIT 1',(amount,))

    def test_backup_restore_and_audit(self):
        self.post('fd_received',1000)
        path=Path(self.temp.name)/'manual.sqlite3'
        self.db.backup(path)
        self.post('fd_received',2000)
        safety=self.db.restore(path)
        self.assertTrue(safety.is_file())
        self.assertEqual(self.l.principal(self.accounts['fd'],'2026-01-31'),1000)
        self.assertEqual(self.db.one("SELECT action FROM audit_log ORDER BY id DESC LIMIT 1")['action'],'restore')
        # Safety backup really contains the replaced records.
        with closing(sqlite3.connect(safety)) as conn: self.assertEqual(conn.execute('SELECT COUNT(*) FROM entries').fetchone()[0],2)
        self.assertEqual(self.db.conn.execute('PRAGMA foreign_keys').fetchone()[0],1)

    def test_backup_validation_rejects_corruption_and_schema_tampering(self):
        bad=Path(self.temp.name)/'bad.sqlite3'; bad.write_bytes(b'not sqlite')
        with self.assertRaises(ValueError): self.db.restore(bad)
        self.db.backup(Path(self.temp.name)/'tampered.sqlite3')
        with closing(sqlite3.connect(Path(self.temp.name)/'tampered.sqlite3')) as conn:
            conn.execute('DROP TRIGGER entries_no_delete'); conn.commit()
        with self.assertRaisesRegex(ValueError,'schema'): self.db.restore(Path(self.temp.name)/'tampered.sqlite3')
        self.assertEqual(len(self.l.cash_accounts()),2)
        self.assertFalse(list((Path(self.temp.name)/'backups').glob('pre-restore-*')))

    def test_backup_does_not_overwrite_existing_files(self):
        with self.assertRaises(ValueError): self.db.backup(self.db.path)
        path=Path(self.temp.name)/'backup.sqlite3'; self.db.backup(path)
        with self.assertRaises(ValueError): self.db.backup(path)

    def test_failed_backup_is_not_published_or_left_as_daily_file(self):
        path=Path(self.temp.name)/'failed.sqlite3'
        with patch.object(self.db,'validate_backup',side_effect=ValueError('Failed validation')):
            with self.assertRaises(ValueError): self.db.backup(path)
        self.assertFalse(path.exists())
        self.assertEqual(list(Path(self.temp.name).glob('.backup-*')),[])

    def test_daily_retention_and_no_repeated_backup(self):
        folder=Path(self.temp.name)/'backups'; folder.mkdir(exist_ok=True)
        first=date(2026,1,1)
        for i in range(35):
            with patch('finance_ledger.storage.today',return_value=(first+timedelta(days=i)).isoformat()): self.db.daily_backup()
        backups=sorted(folder.glob('daily-*.sqlite3'))
        self.assertEqual(len(backups),30)
        latest=backups[-1]; original_time=latest.stat().st_mtime_ns
        with patch('finance_ledger.storage.today',return_value='2026-02-04'): self.db.daily_backup()
        self.assertEqual(latest.stat().st_mtime_ns,original_time)

    def test_single_instance_lock_in_another_process(self):
        code='from finance_ledger.storage import Database, AlreadyRunning\nimport sys\ntry:\n Database(sys.argv[1])\nexcept AlreadyRunning:\n sys.exit(42)\n'
        result=subprocess.run([sys.executable,'-c',code,self.temp.name],capture_output=True,text=True)
        self.assertEqual(result.returncode,42,result.stderr)

    def test_updates_preserve_records_and_backup_before_migration(self):
        self.post('fd_received',1000)
        self.db.close()
        extended=MIGRATIONS+[(2,'CREATE TABLE migration_test(id INTEGER PRIMARY KEY);')]
        with patch('finance_ledger.storage.MIGRATIONS',extended):
            self.db=Database(self.temp.name,daily=False)
            self.assertEqual(self.db.conn.execute('PRAGMA user_version').fetchone()[0],2)
            self.assertEqual(self.db.one('SELECT COUNT(*) n FROM entries')['n'],1)
            backups=list((Path(self.temp.name)/'backups').glob('pre-migration-*'))
            self.assertEqual(len(backups),1)
            with closing(sqlite3.connect(backups[0])) as conn:
                self.assertEqual(conn.execute('PRAGMA user_version').fetchone()[0],1)
                self.assertEqual(conn.execute('SELECT COUNT(*) FROM entries').fetchone()[0],1)

    def test_failed_migration_rolls_back_and_releases_lock(self):
        self.db.close()
        with patch('finance_ledger.storage.MIGRATIONS',MIGRATIONS+[(2,'CREATE TABLE must_rollback(id INTEGER); SELECT nonexistent FROM nonexistent;')]):
            with self.assertRaises(sqlite3.Error): Database(self.temp.name,daily=False)
        self.db=Database(self.temp.name,daily=False)
        self.assertEqual(self.db.conn.execute('PRAGMA user_version').fetchone()[0],1)
        self.assertIsNone(self.db.one("SELECT name FROM sqlite_master WHERE name='must_rollback'"))

    def test_newer_database_is_not_modified(self):
        self.db.conn.execute('PRAGMA user_version=999')
        self.db.close()
        with self.assertRaisesRegex(ValueError,'newer version'): Database(self.temp.name,daily=False)
        with closing(sqlite3.connect(Path(self.temp.name)/'ledger.sqlite3')) as conn: conn.execute('PRAGMA user_version=1')
        self.db=Database(self.temp.name,daily=False)


class MoneyAndExportCase(unittest.TestCase):
    def test_indian_format_and_paise(self):
        self.assertEqual(money(paise('1,23,45,678.90')),'₹1,23,45,678.90')
        self.assertEqual(money(-100005),'-₹1,000.05')
        self.assertEqual(paise('0.10')+paise('0.20'),30)
        for bad in ('1.001','NaN','Infinity','not money'):
            with self.assertRaises(ValueError): paise(bad)
        self.assertEqual(interest_paise(36500,'10'),10)

    def test_csv_safe_from_formulas_and_decimal_amounts(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'report.csv'
            export_report(Report('Test','2026',['Party','Amount'],[['=HYPERLINK("evil")',12345],['Name',-200]],{1}),path)
            with path.open(encoding='utf-8-sig',newline='') as stream: rows=list(csv.reader(stream))
            self.assertTrue(rows[3][0].startswith("'="))
            self.assertEqual(rows[3][1],'123.45')
            self.assertEqual(rows[4][1],'-2.00')

    def test_pdf_and_excel_exports(self):
        with tempfile.TemporaryDirectory() as directory:
            report=Report('Interest received','2026-01-01 to 2026-01-31',['Party','Amount'],[['Demo Party',12345]],{1},['Actual interest only.'])
            for suffix in ('.xlsx','.pdf'):
                path=Path(directory)/('report'+suffix)
                export_report(report,path)
                self.assertGreater(path.stat().st_size,100)
            from openpyxl import load_workbook
            workbook=load_workbook(Path(directory)/'report.xlsx')
            self.assertEqual(workbook.active.cell(5,2).value,123.45)
            self.assertEqual((Path(directory)/'report.pdf').read_bytes()[:4],b'%PDF')

    def test_failed_export_preserves_previous_file(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'report.pdf'; path.write_bytes(b'previous report')
            with patch('finance_ledger.reports._export_report',side_effect=OSError('Disk error')):
                with self.assertRaises(OSError):
                    export_report(Report('Test','2026',['Amount'],[[100]],{0}),path)
            self.assertEqual(path.read_bytes(),b'previous report')
            self.assertEqual(list(Path(directory).glob('.finance-ledger-*')),[])

    def test_demo_is_explicit_and_repeatable(self):
        with tempfile.TemporaryDirectory() as directory:
            db=Database(directory,daily=False)
            self.assertFalse(db.one('SELECT id FROM parties LIMIT 1'))
            seed_demo(db); seed_demo(db)
            self.assertEqual(db.one('SELECT COUNT(*) n FROM entries')['n'],5)
            self.assertEqual(Ledger(db).dashboard('2026-10-01','2026-10-31')['cash'],13_520_000)
            db.close()


if __name__=='__main__':
    unittest.main()
