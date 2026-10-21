"""Local historical reports; estimates never enter actual-interest reports."""
import csv
import os
import tempfile
from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from .ledger import MODULES, TYPES
from .categories import CATEGORY_REPORTS, CATEGORY_TYPES, transaction_detail
from .money import money, valid_date

CASHBOOK_REPORT = 'Cash book (all accounts)'
REPORT_NAMES = (
 CASHBOOK_REPORT,
 'Overall Amount In / Amount Out', 'Cash book / Bank ledger', 'FD summary', 'Individual FD statement',
 'Loan summary', 'Individual loan statement', 'External Loan summary', 'External Loan statement',
 'Hand Loan / Finance Hand Loan', 'Interest received', 'Interest paid by category', 'Party statement',
 'Principal outstanding as of date', 'Monthly summary', 'General income / expense',
) + tuple(CATEGORY_REPORTS.values())


@dataclass
class Report:
    title: str
    period: str
    columns: list
    rows: list
    money_columns: set = field(default_factory=set)
    notes: list = field(default_factory=list)

    def display_rows(self, symbol='₹'):
        return [[money(value,symbol) if i in self.money_columns and isinstance(value,int) else str(value if value is not None else '') for i,value in enumerate(row)] for row in self.rows]


def previous(day):
    return (date.fromisoformat(day)-timedelta(days=1)).isoformat() if day > '0001-01-01' else None


def cashbook_balances(ledger, start, end, cash_id=None):
    """Actual running cash, independent of list search/category/sort/pagination."""
    accounts = [a for a in ledger.cash_accounts() if not cash_id or a['id'] == cash_id]
    before = previous(start)
    opening = sum(ledger.cash_balance(a['id'], before) for a in accounts) if before else 0
    rows = ledger.entries(start, end, cash_id=cash_id)
    events = [(r['tx_date'], r['created_at'], r['id'], r) for r in rows]
    events += [(a['opening_date'], '', 0, a['opening_paise']) for a in accounts if start <= a['opening_date'] <= end]
    balances, balance = {}, opening
    for _, _, _, event in sorted(events, key=lambda e: e[:3]):
        if isinstance(event, int):
            balance += event
        else:
            if cash_id or not event['type'].startswith('transfer'):
                balance += event['amount_paise'] * event['factor'] * event['direction']
            balances[event['id']] = balance
    return opening, balance, balances


def build_cashbook_report(ledger, name, start, end, **filters):
    category = next((key for key, label in CATEGORY_REPORTS.items() if label == name), None)
    filters = dict(filters)
    if category:
        # All transaction shows the family; a selected type narrows its entries.
        filters.pop('module', None)
    def entries(first, last):
        result = ledger.entries(first, last, **filters)
        return [r for r in result if r['type'] in CATEGORY_TYPES[category]] if category else result
    selected = entries(start, end)
    cash_id = filters.get('cash_id')
    complete = not category and not any(filters.get(k) for k in ('party_id', 'finance_id', 'type', 'category_id', 'module', 'status'))
    if not cash_id and not (category == 'transfers'):
        selected = [r for r in selected if not r['type'].startswith('transfer')]
    before = previous(start)
    notes = ['Category means Transaction Details. Cash In / Cash Out are derived from the category. Principal and actual interest remain separate.']
    if any(r['module'] and not r['finance_id'] for r in selected):
        notes.append('Finance entries without linked account history are included in cash movements; they do not create opening principal or change linked-account outstanding balances.')
    if complete:
        opening, closing, balances = cashbook_balances(ledger, start, end, cash_id)
        accounts = [a for a in ledger.cash_accounts() if not cash_id or a['id'] == cash_id]
        introductions = [(a['opening_date'], '', 0, a) for a in accounts if start <= a['opening_date'] <= end]
        balance_label = 'Cash balance'
        notes.append('Actual cash/bank balance. Opening introductions are non-cash. Internal transfers are excluded across all accounts and included for an individual payment account.')
    else:
        history = entries('0001-01-01', before) if before else []
        if not (category == 'transfers'):
            history = [r for r in history if not r['type'].startswith('transfer')]
        opening = sum(r['amount_paise'] * r['factor'] * r['direction'] for r in history)
        closing, balances, introductions = opening, {}, []
        balance_label = 'Category balance'
        notes.append('Opening/closing are cumulative net movements for the selected category/filters, not actual cash or principal outstanding. Use the complete Cash Book for actual cash and account statements for principal.')
    rows = [[start, '', '', 'Opening balance', 0, 0, opening]]
    balance = opening
    events = [(r['tx_date'], r['created_at'], r['id'], r) for r in selected] + introductions
    for day, _, _, item in sorted(events, key=lambda e: e[:3]):
        if 'group_id' not in item:
            balance += item['opening_paise']
            rows.append([day, '', '', 'Opening introduced (non-cash)', 0, 0, balance])
            continue
        amount = item['amount_paise'] * item['factor']
        balance = balances[item['id']] if complete else balance + amount * item['direction']
        rows.append([day, item['party_name'] or item['particulars'], item['place_name'],
                     ('Reversal: ' if item['factor'] == -1 else '') + transaction_detail(item['type']),
                     amount if item['direction'] == 1 else 0, amount if item['direction'] == -1 else 0, balance])
    if not complete:
        closing = balance
    rows.append([end, '', '', 'Closing / totals', sum(r[4] for r in rows), sum(r[5] for r in rows), closing])
    interest_in=sum(r['amount_paise']*r['factor'] for r in selected if r['interest'] and r['direction']==1)
    interest_out=sum(r['amount_paise']*r['factor'] for r in selected if r['interest'] and r['direction']==-1)
    notes.append(f'Actual interest received: {money(interest_in)}. Actual interest paid: {money(interest_out)}. These amounts are included in Amount In / Amount Out. Interest and principal entries are identified separately in Transaction Details.')
    return Report(name, f'{start} to {end} · INR · Asia/Kolkata',
                  ['Date', 'Name', 'Place Name', 'Transaction Details', 'Amount In', 'Amount Out', balance_label],
                  rows, {4, 5, 6}, notes)


def build_report(ledger, name, start, end, **filters):
    if name not in REPORT_NAMES:
        raise ValueError('Select a report.')
    valid_date(start); valid_date(end)
    if start > end:
        raise ValueError('Start date must be on or before end date.')
    if name == CASHBOOK_REPORT or name in CATEGORY_REPORTS.values():
        return build_cashbook_report(ledger, name, start, end, **filters)
    period = f'{start} to {end} · INR · Asia/Kolkata'
    before = previous(start)
    account_id = filters.get('finance_id')
    cash_id = filters.get('cash_id')
    party_id = filters.get('party_id')
    notes = ['Principal movements are separate from income and expenses. Reversals reduce the original Amount In/Out on the reversal date.', 'Actual posted amounts only; interest estimates are excluded.']
    if name in ('FD summary','Loan summary','External Loan summary','Hand Loan / Finance Hand Loan','Principal outstanding as of date'):
        modules = {'FD summary':{'fd'},'Loan summary':{'loan'},'External Loan summary':{'other'},'Hand Loan / Finance Hand Loan':{'hand_given','hand_taken'}}.get(name,set(MODULES))
        rows = []
        for a in ledger.finance_accounts():
            if a['start_date'] > end or a['module'] not in modules or (party_id and a['party_id']!=party_id) or (account_id and a['id']!=account_id) or (filters.get('status') and a['status']!=filters['status']) or (filters.get('module') and a['module']!=filters['module']):
                continue
            entries = ledger.entries(start,end,**(filters | {'finance_id':a['id']}))
            amount_in = sum(r['amount_paise']*r['factor'] for r in entries if r['direction']==1)
            amount_out = sum(r['amount_paise']*r['factor'] for r in entries if r['direction']==-1)
            opening = ledger.principal(a['id'],before) if before else 0
            introduced = a['opening_paise'] if start <= a['start_date'] <= end else 0
            actual = sum(r['amount_paise']*r['factor']*r['direction'] for r in entries if r['interest'])
            rows.append([a['reference'],a['party_name'],MODULES[a['module']], 'Receivable' if a['module'] in ('loan','hand_given') else 'Liability',a['status'],opening,introduced,amount_in,amount_out,ledger.principal(a['id'],end),actual])
        details=list(rows)
        for position in ('Receivable','Liability'):
            selected=[row for row in details if row[3]==position]
            if selected:
                rows.append(['TOTAL','','',position,'']+[sum(row[i] for row in selected) for i in range(5,11)])
        return Report(name,period,['Account','Party','Module','Position','Status','Opening principal','Opening introduced','Amount In','Amount Out','Closing principal','Net actual interest'],rows,{5,6,7,8,9,10},notes+['Closing principal is reconstructed as of the end date. Opening introduced is a non-cash brought-forward balance. Movement filters apply to movements only; principal balances always include all principal entries. Status is the currently recorded account status. Receivable and liability totals remain separate.'])
    if 'statement' in name and name != 'Party statement':
        module = {'Individual FD statement':'fd','Individual loan statement':'loan','External Loan statement':'other'}[name]
        if not account_id:
            raise ValueError('Select a finance account for this statement.')
        a = ledger.db.one('SELECT * FROM finance_accounts WHERE id=?', (account_id,))
        if not a or a['module'] != module:
            raise ValueError('Select an account matching the statement module.')
        opening = ledger.principal(account_id,before) if before else 0
        rows = [[start,'','Opening principal','',0,0,opening]]
        balance = opening
        if start <= a['start_date'] <= end and a['opening_paise']:
            rows.append([a['start_date'],'','Opening principal introduced (non-cash)','',0,0,a['opening_paise']])
        movements = ledger.entries(start,end,finance_id=account_id)
        events = [(r['tx_date'],r['created_at'],r) for r in movements]
        if start <= a['start_date'] <= end and a['opening_paise']:
            events.append((a['start_date'],'',None))
            rows = rows[:1]
        for day,_,r in sorted(events,key=lambda e:(e[0],e[1],e[2]['id'] if e[2] else 0)):
            if r is None:
                balance += a['opening_paise']
                rows.append([day,'','Opening principal introduced (non-cash)','',0,0,balance,'']); continue
            signed = r['amount_paise']*r['factor']
            balance += signed*r['principal_delta']
            rows.append([day,r['reference'],('Reversal: ' if r['factor']==-1 else '')+TYPES[r['type']].label,r['cash_name'],signed if r['direction']==1 else 0,signed if r['direction']==-1 else 0,balance])
        rows.append([end,'','Closing / totals','',sum(r[4] for r in rows),sum(r[5] for r in rows),balance])
        return Report(name+' · '+a['reference'],period,['Date','Reference','Type','Payment account','Amount In','Amount Out','Principal balance'],rows,{4,5,6},notes)
    entries = ledger.entries(start,end,**filters)
    if name == 'Interest received':
        entries = [r for r in entries if r['interest'] and r['direction']==1]
    elif name == 'Interest paid by category':
        entries = [r for r in entries if r['interest'] and r['direction']==-1]
    elif name == 'General income / expense':
        entries = [r for r in entries if r['type'].startswith('general')]
    if name == 'Party statement' and not party_id:
        raise ValueError('Select a party for a party statement.')
    if name == 'Cash book / Bank ledger' and not cash_id:
        raise ValueError('Select a cash or bank account for its ledger.')
    cash_ledger = name == 'Cash book / Bank ledger'
    full_cash_scope = name in ('Overall Amount In / Amount Out','Monthly summary') and not any(filters.get(key) for key in ('party_id','finance_id','module','status','type','category_id'))
    scope_accounts = [a for a in ledger.cash_accounts() if not cash_id or a['id']==cash_id]
    if not cash_ledger:
        entries = [r for r in entries if not r['type'].startswith('transfer')]
    if name == 'Monthly summary':
        totals = {}
        for r in entries:
            m = r['tx_date'][:7]
            values = totals.setdefault(m,[0,0,0,0])
            signed = r['amount_paise']*r['factor']
            values[0 if r['direction']==1 else 1] += signed
            if r['interest']:
                values[2 if r['direction']==1 else 3] += signed
        # Include quiet months too.
        cursor = date.fromisoformat(start).replace(day=1)
        while cursor.isoformat()[:7] <= end[:7]:
            totals.setdefault(cursor.isoformat()[:7],[0,0,0,0])
            cursor = date(cursor.year+1,1,1) if cursor.month==12 else date(cursor.year,cursor.month+1,1)
        rows=[]
        for month,v in sorted(totals.items()):
            month_start=max(start,month+'-01')
            year,number=map(int,month.split('-'))
            next_month=date(year+1,1,1) if number==12 else date(year,number+1,1)
            month_end=min(end,(next_month-timedelta(days=1)).isoformat())
            prev=previous(month_start)
            if full_cash_scope:
                opening=sum(ledger.cash_balance(a['id'],prev) for a in scope_accounts) if prev else 0
                introduced=sum(a['opening_paise'] for a in scope_accounts if month_start<=a['opening_date']<=month_end)
                closing=sum(ledger.cash_balance(a['id'],month_end) for a in scope_accounts)
            else:
                historic=ledger.entries('0001-01-01',prev,**filters) if prev else []
                opening=sum(r['amount_paise']*r['factor']*r['direction'] for r in historic if not r['type'].startswith('transfer'))
                introduced=0
                closing=opening+v[0]-v[1]
            rows.append([month,opening,introduced,v[0],v[1],v[0]-v[1],closing,v[2],v[3],v[2]-v[3]])
        return Report(name,period,['Month','Opening balance','Non-cash opening introduced','Amount In','Amount Out','Net cash movement','Closing balance','Interest received','Interest paid','Net interest income'],rows,set(range(1,10)),notes+['Net cash movement is not profit. Internal transfers are excluded from Amount In / Out.','Opening / closing are cash balances for the complete scope; when filtered by party, module, status or type they represent cumulative filtered movement. With an individual payment account, balances also reflect transfers.'])
    if full_cash_scope:
        opening=sum(ledger.cash_balance(a['id'],before) for a in scope_accounts) if before else 0
        introductions=[(a['opening_date'],a['opening_paise']) for a in scope_accounts if start<=a['opening_date']<=end]
        if cash_id:
            entries=ledger.entries(start,end,**filters)
        notes.append('Opening and closing include payment account opening balances. Non-cash opening introductions are listed separately. Internal transfers are excluded from overall totals; for one payment account they appear in its running balance.')
    elif cash_ledger:
        cash = ledger.db.one('SELECT * FROM cash_accounts WHERE id=?', (cash_id,))
        opening = ledger.cash_balance(cash_id,before) if before else 0
        introductions = [(cash['opening_date'],cash['opening_paise'])] if start <= cash['opening_date'] <= end else []
        notes.append('Running balance is the payment account cash balance. Internal transfers are included in this individual ledger. Movement filters do not alter its opening balance.')
    else:
        historic = ledger.entries('0001-01-01',before,**filters) if before else []
        if name == 'Interest received': historic = [r for r in historic if r['interest'] and r['direction']==1]
        if name == 'Interest paid by category': historic = [r for r in historic if r['interest'] and r['direction']==-1]
        if name == 'General income / expense': historic = [r for r in historic if r['type'].startswith('general')]
        opening = sum(r['amount_paise']*r['direction']*r['factor'] for r in historic if not r['type'].startswith('transfer'))
        introductions = []
        notes.append('Running balance shows cumulative net movement for the selected report; it is not profit or principal outstanding.')
    rows = [[start,'','Opening balance','','','',0,0,opening,'']]
    balance = opening
    events = [(r['tx_date'],r['created_at'],r) for r in entries]+[(d,'',amount) for d,amount in introductions]
    for day,_,r in sorted(events,key=lambda e:(e[0],e[1],e[2]['id'] if isinstance(e[2],dict) else 0)):
        if isinstance(r,int):
            balance += r; rows.append([day,'','Opening balance introduced (non-cash)','','','',0,0,balance,'']); continue
        amount = r['amount_paise']*r['factor']
        balance += amount*r['direction']
        rows.append([day,r['reference'],('Reversal: ' if r['factor']==-1 else '')+TYPES[r['type']].label,r['party_name'] or r['particulars'],r['cash_name'],r['category_name'] or MODULES.get(r['module'],''),amount if r['direction']==1 else 0,amount if r['direction']==-1 else 0,balance,r['place_name']])
    rows.append([end,'','Closing / totals','','','',sum(r[6] for r in rows if 'Internal transfer' not in str(r[2]) or cash_ledger),sum(r[7] for r in rows if 'Internal transfer' not in str(r[2]) or cash_ledger),balance,''])
    if name == 'Party statement':
        accounts = [a for a in ledger.finance_accounts() if a['party_id']==party_id and a['start_date']<=end]
        for a in accounts:
            notes.append(f'{a["reference"]} · {MODULES[a["module"]]}: principal outstanding {money(ledger.principal(a["id"],end),"Rs ")}.')
        principal_in = sum(r['amount_paise']*r['factor'] for r in entries if r['principal_delta'] and r['direction']==1)
        principal_out = sum(r['amount_paise']*r['factor'] for r in entries if r['principal_delta'] and r['direction']==-1)
        interest_in = sum(r['amount_paise']*r['factor'] for r in entries if r['interest'] and r['direction']==1)
        interest_out = sum(r['amount_paise']*r['factor'] for r in entries if r['interest'] and r['direction']==-1)
        notes.append(f'Period principal In / Out: {money(principal_in,"Rs ")} / {money(principal_out,"Rs ")}. Actual interest In / Out: {money(interest_in,"Rs ")} / {money(interest_out,"Rs ")}.')
    if name == 'Interest paid by category':
        for module in ('fd','other','hand_taken'):
            total = sum(r['amount_paise']*r['factor'] for r in entries if r['module']==module)
            notes.append(f'{MODULES[module]} actual interest paid: {money(total,"Rs ")}')
    return Report(name,period,['Date','Name','Place Name','Transaction Details','Amount In','Amount Out','Running balance'],
                  [[r[i] for i in (0,3,9,2,6,7,8)] for r in rows],{4,5,6},notes)


def safe_text(value):
    text = str(value)
    return "'"+text if text.lstrip().startswith(('=','+','-','@','\t','\r')) else text


def export_report(report, path, business_name='Finance Ledger'):
    path=Path(path)
    # A failed export must not truncate a previously saved report.
    handle,temporary=tempfile.mkstemp(prefix='.finance-ledger-',suffix=path.suffix,dir=path.parent)
    os.close(handle)
    try:
        _export_report(report,Path(temporary),business_name)
        os.replace(temporary,path)
    finally:
        Path(temporary).unlink(missing_ok=True)
    return path


def _export_report(report, path, business_name='Finance Ledger'):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == '.csv':
        with path.open('w',newline='',encoding='utf-8-sig') as stream:
            writer = csv.writer(stream)
            writer.writerow([safe_text(business_name),safe_text(report.title)])
            writer.writerow([report.period]); writer.writerow(report.columns)
            for row in report.rows:
                writer.writerow([f'{Decimal(value)/100:.2f}' if i in report.money_columns and isinstance(value,int) else safe_text(value or '') for i,value in enumerate(row)])
            for note in report.notes:
                writer.writerow([safe_text(note)])
    elif suffix == '.xlsx':
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise ValueError('Excel export needs openpyxl. Install requirements.txt or use CSV.') from None
        workbook = Workbook(); sheet = workbook.active; sheet.title = 'Finance Ledger'
        sheet.append([safe_text(business_name)]); sheet.append([safe_text(report.title)]); sheet.append([report.period]); sheet.append(report.columns)
        for row in report.rows:
            values=[]
            for i,value in enumerate(row):
                if i in report.money_columns and isinstance(value,int): value=Decimal(value)/100
                elif report.columns[i]=='Date' and value:
                    value=date.fromisoformat(value)
                else: value=safe_text(value if value is not None else '')
                values.append(value)
            sheet.append(values)
        last_column=get_column_letter(len(report.columns)); last_row=sheet.max_row
        sheet.sheet_view.showGridLines=False
        for number in (1,2,3):
            if len(report.columns)>1: sheet.merge_cells(f'A{number}:{last_column}{number}')
            sheet.cell(number,1).font=Font(name='Calibri',size=16 if number==1 else 11,bold=number<3,color='164E45')
            sheet.row_dimensions[number].height=26 if number==1 else 22
        edge=Side(style='thin',color='CED8D2')
        border=Border(left=edge,right=edge,top=edge,bottom=edge)
        for row in sheet.iter_rows(min_row=4,max_row=last_row):
            for i,cell in enumerate(row):
                cell.font=Font(name='Calibri',size=11)
                cell.border=border
                cell.alignment=Alignment(vertical='center',wrap_text=True,horizontal='right' if i in report.money_columns else 'left')
                if i in report.money_columns:
                    cell.number_format='#,##0.00;[Red]-#,##0.00;"–"'
                if report.columns[i]=='Date' and cell.row>4: cell.number_format='dd.mm.yyyy'
            sheet.row_dimensions[row[0].row].height=32
        for cell in sheet[4]:
            cell.font=Font(name='Calibri',size=11,color='FFFFFF',bold=True)
            cell.fill=PatternFill('solid',fgColor='164E45')
        sheet.row_dimensions[4].height=38
        is_cashbook='Transaction Details' in report.columns
        if is_cashbook:
            for number in (5,last_row):
                for cell in sheet[number]:
                    cell.font=Font(name='Calibri',size=11,bold=True,color='164E45')
                    cell.fill=PatternFill('solid',fgColor='E8F0EB')
        sheet.freeze_panes='E6' if is_cashbook else 'A5'
        sheet.auto_filter.ref=f'A4:{last_column}{max(4,last_row-1) if is_cashbook else last_row}'
        for i,column in enumerate(report.columns,1):
            width=12 if column=='S.No' else 15 if column=='Date' else 36 if column in ('Transaction Details','Transaction Details / Category','Particulars / Description') else 25 if column=='Name' else min(32,max(20,len(column)+2))
            sheet.column_dimensions[get_column_letter(i)].width=width
        sheet.append([])
        for note in report.notes:
            sheet.append([safe_text(note)])
            number=sheet.max_row
            if len(report.columns)>1: sheet.merge_cells(f'A{number}:{last_column}{number}')
            sheet.cell(number,1).alignment=Alignment(wrap_text=True,vertical='top')
            sheet.cell(number,1).font=Font(name='Calibri',size=10,color='53685D')
            sheet.row_dimensions[number].height=30
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'; sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = '1:4'
        sheet.print_options.horizontalCentered=True
        sheet.print_area=f'A1:{last_column}{sheet.max_row}'
        workbook.save(path)
    elif suffix == '.pdf':
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
            from xml.sax.saxutils import escape
        except ImportError:
            raise ValueError('PDF export needs reportlab. Install requirements.txt or use CSV.') from None
        styles = getSampleStyleSheet()
        small = ParagraphStyle('LedgerSmall',fontSize=7,leading=9)
        def para(text,style=small):
            return Paragraph(escape(str(text).replace('₹','Rs ').replace('·',' / ').replace('→',' to ')),style)
        doc = SimpleDocTemplate(str(path),pagesize=landscape(A4),rightMargin=24,leftMargin=24,topMargin=26,bottomMargin=30)
        content = [para(business_name,styles['Title']),para(report.title,styles['Heading2']),para(report.period),Spacer(1,12)]
        data = [[para(c) for c in report.columns]]+[[para(c) for c in row] for row in report.display_rows('Rs ')]
        weights = [1.0]*len(report.columns)
        for i,column in enumerate(report.columns):
            if column in ('Transaction type','Transaction Details','Name','Place Name','Type','Party','Module','Category'): weights[i] = 1.7
        widths = [w/sum(weights)*(landscape(A4)[0]-48) for w in weights]
        table = LongTable(data,colWidths=widths,repeatRows=1,hAlign='LEFT')
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dcece6')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f4f7f6')]),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),('LINEBELOW',(0,0),(-1,0),0.5,colors.HexColor('#79968c'))]))
        content.append(table); content.append(Spacer(1,12))
        content.extend(para(note) for note in report.notes)
        def footer(canvas,document):
            canvas.setFont('Helvetica',8); canvas.drawRightString(landscape(A4)[0]-24,16,f'Finance Ledger | Page {document.page}')
        doc.build(content,onFirstPage=footer,onLaterPages=footer)
    else:
        raise ValueError('Choose a .csv, .xlsx or .pdf file.')
    return path
